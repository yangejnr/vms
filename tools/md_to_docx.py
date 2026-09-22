#!/usr/bin/env python3
"""Convert the project's Markdown documentation into Word (.docx) files.

Usage:
    python tools/md_to_docx.py                     # convert all docs
    python tools/md_to_docx.py docs/USER_GUIDE.md  # convert one file

Supports the subset of Markdown used by this project: headings, paragraphs,
bold/italic/code spans, links, images, tables, ordered and unordered lists,
fenced code blocks, blockquotes and horizontal rules. Headings use Word's
built-in Heading styles so the navigation pane and an automatic table of
contents both work.

Requires: python-docx  (pip install python-docx)
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ----------------------------------------------------------------------------- config

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(PROJECT_ROOT, 'docs')

NCS_GREEN = RGBColor(0x0F, 0x6F, 0x3F)
NCS_DEEP = RGBColor(0x12, 0x36, 0x29)
MUTED = RGBColor(0x5E, 0x6F, 0x66)
CODE_SHADE = 'F4F7F5'

# Order matters: the index is converted last.
DOC_FILES = [
    'USER_GUIDE.md',
    'MANAGEMENT_PRESENTATION.md',
    'TECHNICAL_DOCUMENTATION.md',
    'BUSINESS_REQUIREMENTS.md',
    'README.md',
]


# ----------------------------------------------------------------------- helpers

def shade(element, fill):
    """Apply a background shade to a paragraph."""
    pr = element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def shade_run(run, fill):
    """Apply a background shade to a single text run (used for inline code)."""
    r_pr = run._element.get_or_add_rPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    r_pr.append(shd)


def set_cell_background(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_hyperlink(paragraph, url, text):
    """Add a clickable hyperlink run to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0F6F3F')
    r_pr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement('w:t')
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_code_paragraph(doc, text, first=False, last=False):
    """Add a monospace paragraph with shading to look like a code block."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0 if not first else 6)
    pf.space_after = Pt(0 if not last else 6)
    pf.left_indent = Inches(0.18)
    run = p.add_run(text if text else ' ')
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    r = run._element
    r_pr = r.get_or_add_rPr()
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Consolas')
    fonts.set(qn('w:hAnsi'), 'Consolas')
    r_pr.append(fonts)
    shade(p._p, CODE_SHADE)
    return p


INLINE_RE = re.compile(
    r'(\*\*.+?\*\*)'          # bold
    r'|(`[^`]+?`)'            # code
    r'|(\*[^*]+?\*)'          # italic
    r'|(\[[^\]]+?\]\([^)]+?\))'  # link
    r'|(~~.+?~~)'             # strikethrough
)


def add_inline(paragraph, text, base_size=None):
    """Render inline markdown (bold, italic, code, links) into a paragraph."""
    text = text.replace('<br>', '\n').replace('<br/>', '\n')
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            _add_plain(paragraph, text[pos:match.start()], base_size)
        token = match.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt((base_size or 10.5) - 1)
            r_pr = run._element.get_or_add_rPr()
            fonts = OxmlElement('w:rFonts')
            fonts.set(qn('w:ascii'), 'Consolas')
            fonts.set(qn('w:hAnsi'), 'Consolas')
            r_pr.append(fonts)
            shade_run(run, CODE_SHADE)
        elif token.startswith('~~'):
            run = paragraph.add_run(token[2:-2])
            run.font.strike = True
        elif token.startswith('['):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token).groups()
            if url.startswith('http'):
                add_hyperlink(paragraph, url, label)
            else:
                # Internal anchor: render the label only, anchors do not resolve in Word.
                run = paragraph.add_run(label)
                run.font.color.rgb = NCS_GREEN
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        if base_size:
            for run in paragraph.runs:
                if run.font.size is None:
                    run.font.size = Pt(base_size)
        pos = match.end()
    if pos < len(text):
        _add_plain(paragraph, text[pos:], base_size)
    if base_size:
        for run in paragraph.runs:
            if run.font.size is None:
                run.font.size = Pt(base_size)


def _add_plain(paragraph, text, base_size=None):
    run = paragraph.add_run(text)
    if base_size:
        run.font.size = Pt(base_size)
    return run


def is_table_separator(line):
    return bool(re.match(r'^\s*\|?[\s:|-]+\|[\s:|-]*$', line)) and '-' in line


def parse_table_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


# ------------------------------------------------------------------ main converter

def add_field(paragraph, instruction):
    """Insert a Word field code (used for PAGE and NUMPAGES)."""
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '1'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(t)
    run._r.append(fld_end)
    return run


def add_furniture(doc, title):
    """Add a running header and a footer with page numbers."""
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.text = ''
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run(title)
        hr.font.size = Pt(8)
        hr.font.color.rgb = MUTED
        hr.italic = True

        footer = section.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.text = ''
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run(f'{title}  ·  Page ')
        fr.font.size = Pt(8)
        fr.font.color.rgb = MUTED
        pr = add_field(fp, ' PAGE ')
        pr.font.size = Pt(8)
        pr.font.color.rgb = MUTED
        mid = fp.add_run(' of ')
        mid.font.size = Pt(8)
        mid.font.color.rgb = MUTED
        tr = add_field(fp, ' NUMPAGES ')
        tr.font.size = Pt(8)
        tr.font.color.rgb = MUTED


def add_toc_field(doc):
    """Insert an auto-updating Word table of contents field."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    hint = OxmlElement('w:t')
    hint.text = 'Right-click and choose "Update Field" to build the table of contents.'
    run._r.append(hint)
    run._r.append(fld_end)
    note = doc.add_paragraph()
    nr = note.add_run('(In Word: right-click the line above → Update Field → Update entire table.)')
    nr.italic = True
    nr.font.size = Pt(8)
    nr.font.color.rgb = MUTED


def strip_front_matter(lines):
    """Remove leading YAML/Marp front matter delimited by '---' lines."""
    if lines and lines[0].strip() == '---':
        for idx in range(1, len(lines)):
            if lines[idx].strip() in ('---', '...'):
                return lines[idx + 1:]
    return lines


def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    lines = strip_front_matter(lines)

    doc = Document()

    # Base typography
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for level, size, color in [
        (1, 22, NCS_DEEP), (2, 16, NCS_DEEP), (3, 13, NCS_GREEN), (4, 11.5, NCS_GREEN),
    ]:
        style = doc.styles[f'Heading {level}']
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    # Page margins
    for section in doc.sections:
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    # Running header/footer, using the document's first heading as the title.
    doc_title = os.path.splitext(os.path.basename(md_path))[0].replace('_', ' ').title()
    for candidate in lines:
        m = re.match(r'^#\s+(.*)$', candidate.strip())
        if m:
            doc_title = re.sub(r'\{#[^}]+\}', '', m.group(1)).strip()
            break
    add_furniture(doc, doc_title)

    base_dir = os.path.dirname(os.path.abspath(md_path))

    i = 0
    in_code = False
    code_lines = []
    code_lang = ''
    image_count = 0
    table_count = 0
    skipped_images = []
    first_heading_done = False

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # ---------------------------------------------------------- code fences
        if stripped.startswith('```'):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                in_code = False
                for idx, cl in enumerate(code_lines):
                    add_code_paragraph(
                        doc, cl,
                        first=(idx == 0),
                        last=(idx == len(code_lines) - 1),
                    )
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue

        # ------------------------------------------------------- horizontal rule
        if re.match(r'^\s*(---|\*\*\*|___)\s*$', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p_pr = p._p.get_or_add_pPr()
            borders = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), 'DCE5DF')
            borders.append(bottom)
            p_pr.append(borders)
            i += 1
            continue

        # --------------------------------------------------------------- images
        img_match = re.match(r'^\s*!\[([^\]]*)\]\(([^)]+)\)\s*$', line)
        if img_match:
            alt, src = img_match.groups()
            img_path = src if os.path.isabs(src) else os.path.join(base_dir, src)
            if os.path.exists(img_path):
                try:
                    doc.add_picture(img_path, width=Inches(5.9))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if alt:
                        cap = doc.add_paragraph()
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = cap.add_run(alt)
                        run.italic = True
                        run.font.size = Pt(9)
                        run.font.color.rgb = MUTED
                    image_count += 1
                except (UnrecognizedImageError, Exception) as exc:  # noqa: BLE001
                    skipped_images.append(f'{src} ({exc})')
                    _add_plain(doc.add_paragraph(), f'[image: {src}]')
            else:
                skipped_images.append(f'{src} (not found)')
                _add_plain(doc.add_paragraph(), f'[image: {src}]')
            i += 1
            continue

        # --------------------------------------------------------------- tables
        if stripped.startswith('|') and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            header = parse_table_row(line)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                body.append(parse_table_row(lines[i]))
                i += 1

            cols = max([len(header)] + [len(r) for r in body]) if body else len(header)
            table = doc.add_table(rows=1, cols=cols)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            hdr = table.rows[0].cells
            for c in range(cols):
                text = header[c] if c < len(header) else ''
                cell = hdr[c]
                cell.text = ''
                para = cell.paragraphs[0]
                add_inline(para, text, base_size=9.5)
                for run in para.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_background(cell, '0F6F3F')

            for row_data in body:
                cells = table.add_row().cells
                for c in range(cols):
                    text = row_data[c] if c < len(row_data) else ''
                    cells[c].text = ''
                    add_inline(cells[c].paragraphs[0], text, base_size=9.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            table_count += 1
            continue

        # ------------------------------------------------------------- headings
        h = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if h:
            level = len(h.group(1))
            text = h.group(2).strip()
            text = re.sub(r'\{#[^}]+\}', '', text).strip()
            in_toc = bool(re.match(r'^\s*\d+\.\s+\[', text)) and False
            heading = doc.add_heading('', level=min(level, 4))
            add_inline(heading, text)

            # A hand-written "Table of Contents" section becomes an auto-updating
            # Word TOC field: skip the numbered link list that follows it.
            if text.lower().strip() == 'table of contents':
                while i + 1 < len(lines):
                    nxt = lines[i + 1].strip()
                    if not nxt or re.match(r'^\s*[-*]?\s*\d+\.\s*\[', nxt) or nxt.startswith('   - '):
                        i += 1
                        continue
                    break
                add_toc_field(doc)
            i += 1
            continue

        # ----------------------------------------------------------- blockquote
        if stripped.startswith('>'):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote_lines.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            p_pr = p._p.get_or_add_pPr()
            borders = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '18')
            left.set(qn('w:color'), 'D9A441')
            borders.append(left)
            p_pr.append(borders)
            shade(p._p, 'F4F7F5')
            add_inline(p, ' '.join(ql for ql in quote_lines if ql))
            continue

        # ----------------------------------------------------------------- lists
        ol = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        ul = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if ol or ul:
            indent = len(ul.group(1)) if ul else 0
            text = ol.group(2) if ol else ul.group(2)
            style = 'List Number' if ol else 'List Bullet'
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if indent >= 2:
                p.paragraph_format.left_indent = Inches(0.55)
            add_inline(p, text)
            i += 1
            continue

        # ------------------------------------------------------------ empty line
        if not stripped:
            i += 1
            continue

        # ------------------------------------ Marp/YAML front-matter directives
        if stripped in ('---', '...') or re.match(r'^<!--.*-->$', stripped):
            i += 1
            continue
        if re.match(r'^(marp|theme|paginate|size|header|footer|style|class|_class|_paginate|_header|_footer)\s*:', stripped):
            i += 1
            continue
        if stripped.startswith('_class:') or stripped.startswith('_paginate:'):
            i += 1
            continue

        # ------------------------------------------------------------- paragraph
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith('#') or nxt.startswith('```')
                    or nxt.startswith('>') or nxt.startswith('|')
                    or re.match(r'^\s*!\[', lines[i])
                    or re.match(r'^\s*(\d+)\.\s', lines[i])
                    or re.match(r'^\s*[-*+]\s', lines[i])
                    or re.match(r'^\s*(---|\*\*\*|___)\s*$', lines[i])):
                break
            para_lines.append(nxt)
            i += 1
        add_inline(doc.add_paragraph(), ' '.join(para_lines))

    doc.save(docx_path)
    return {
        'images': image_count,
        'tables': table_count,
        'skipped': skipped_images,
    }


def main():
    args = sys.argv[1:]
    targets = args if args else [os.path.join(DOCS_DIR, f) for f in DOC_FILES]

    print(f'Converting {len(targets)} document(s)\n')
    ok = True
    for md_path in targets:
        if not os.path.exists(md_path):
            print(f'  SKIP  {md_path} (not found)')
            ok = False
            continue
        docx_path = os.path.splitext(md_path)[0] + '.docx'
        try:
            stats = convert(md_path, docx_path)
            size = os.path.getsize(docx_path) / 1024
            print(f'  OK    {os.path.basename(docx_path):36} '
                  f'{size:7.1f} KB  images={stats["images"]:2}  tables={stats["tables"]:2}')
            if stats['skipped']:
                for s in stats['skipped']:
                    print(f'          ! skipped image: {s}')
        except Exception as exc:  # noqa: BLE001
            print(f'  FAIL  {md_path}: {exc}')
            ok = False
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())